"""
Builds the downloadable Excel version of the Attrition trend chart - a real,
editable Excel chart (not a picture of one), so HR can tweak a number that
looks off in the data table and have the chart reflect it, then drop the
whole thing into a report. See app.py's /api/attrition/trends.xlsx.

Mirrors client/src/components/AttritionTrendChart.jsx: same columns/order
(Opening, Exits, Joins, Closing, Attrition rate), same series colors, and
the same TWO separate single-axis bar charts rather than one - Opening/
Closing headcount run in the thousands while Exits/Joins run in the tens-
to-low-hundreds, so on one shared axis the Exits/Joins bars were unreadably
tiny next to headcount (an earlier dual-axis version fixed that but was
flagged as too complex to read at a glance; splitting by scale keeps each
chart a plain single-axis bar chart while both pairs stay readable sizes).
"""
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

# Same hex values as client/src/index.css's --accent/--late/--ontime/--accent-2.
OPENING_COLOR = "42145F"
EXITS_COLOR = "E34948"
JOINS_COLOR = "008300"
CLOSING_COLOR = "00B2A9"

COLUMNS = ["Period", "Opening headcount", "Exits", "Joins", "Closing headcount", "Attrition rate (%)"]

MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def format_period_label(period):
    """Mirrors client/src/lib/format.js's formatPeriodLabel - "YYYY-MM" ->
    "Jan '26", "YYYY-Qn" -> "Qn '26", "YYYY" stays as-is."""
    m = re.match(r"^(\d{4})-Q(\d)$", period)
    if m:
        return f"Q{m.group(2)} '{m.group(1)[2:]}"
    m = re.match(r"^(\d{4})-(\d{2})$", period)
    if m:
        return f"{MONTH_NAMES[int(m.group(2)) - 1]} '{m.group(1)[2:]}"
    return period


def build_trend_workbook(periods):
    """periods: list of dicts from db.query_attrition_trends (period,
    headcount_start, exits, joins, headcount_end, attrition_rate). Returns
    the .xlsx file's bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attrition trend"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for p in periods:
        ws.append([
            format_period_label(p["period"]),
            p["headcount_start"], p["exits"], p["joins"], p["headcount_end"], p["attrition_rate"],
        ])

    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0) + 2
        ws.column_dimensions[col_cells[0].column_letter].width = width

    n = len(periods)
    if n > 0:
        last_row = n + 1
        cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)

        def make_chart(title, min_col, max_col, colors):
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
            chart.title = title
            chart.style = 2
            chart.height = 9
            chart.width = 26
            chart.add_data(Reference(ws, min_col=min_col, max_col=max_col, min_row=1, max_row=last_row), titles_from_data=True)
            chart.set_categories(cats)
            for series, color in zip(chart.series, colors):
                series.graphicalProperties.solidFill = color
            return chart

        # Columns B,E = Opening/Closing headcount; C,D = Exits/Joins - two
        # charts, not one combined, same reasoning as the dashboard's split.
        headcount_chart = make_chart("Headcount", 2, 2, (OPENING_COLOR,))
        headcount_chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=last_row), titles_from_data=True)
        headcount_chart.series[-1].graphicalProperties.solidFill = CLOSING_COLOR

        flow_chart = make_chart("Exits & Joins", 3, 4, (EXITS_COLOR, JOINS_COLOR))

        ws.add_chart(headcount_chart, f"A{last_row + 3}")
        ws.add_chart(flow_chart, f"A{last_row + 21}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
